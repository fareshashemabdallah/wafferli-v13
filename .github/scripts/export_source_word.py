#!/usr/bin/env python3
"""Export every tracked text source file as a navigable, searchable Word reference."""
import os
import subprocess
from datetime import datetime
from pathlib import Path
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT=Path("Dabberna_Complete_Source_Code.docx")
BINARY={".png",".jpg",".jpeg",".gif",".ico",".webp",".pdf",".woff",".woff2",".ttf",".otf",".zip",".docx",".mp4",".svg"}
GREEN=RGBColor(38,62,51)
MUTED=RGBColor(102,120,110)

def kind(name):
    if name.startswith("app/src/routes/"):return "01 — Pages and routes"
    if name.startswith("app/src/components/"):return "02 — Components and UI"
    if name.startswith("app/src/lib/"):return "03 — Logic, state and tests"
    if name.startswith("app/src/"):return "04 — App entry and styling"
    if name.startswith("app/scripts/"):return "05 — Supporting scripts"
    if name.startswith(".github/"):return "06 — Automation and deployment"
    if name.startswith("app/"):return "07 — App configuration and lockfiles"
    return "08 — Root files and documentation"

def set_font(style,name,size,color):
    style.font.name=name
    style.font.size=Pt(size)
    style.font.color.rgb=color
    rpr=style.element.get_or_add_rPr()
    fonts=rpr.rFonts
    if fonts is None:
        fonts=OxmlElement("w:rFonts");rpr.insert(0,fonts)
    fonts.set(qn("w:ascii"),name)
    fonts.set(qn("w:hAnsi"),name)

def page_number(para):
    field=OxmlElement("w:fldSimple")
    field.set(qn("w:instr"),"PAGE")
    para._p.append(field)

def source_files():
    paths=[x.decode() for x in subprocess.check_output(["git","ls-files","-z"]).split(b"\0") if x]
    valid=[];excluded=[]
    for name in paths:
        f=Path(name)
        if not f.is_file():continue
        if f.suffix.lower() in BINARY:
            excluded.append(name);continue
        try:data=f.read_text(encoding="utf-8")
        except (UnicodeError,OSError):
            excluded.append(name);continue
        if "\0" in data:
            excluded.append(name);continue
        valid.append((name,data))
    valid.sort(key=lambda item:(kind(item[0]),item[0].lower()))
    return valid,excluded

def main():
    files,excluded=source_files()
    assert len(files)>100, "Expected the full Dabberna repository, not a partial checkout"
    doc=Document()
    section=doc.sections[0]
    section.page_width=Inches(8.27);section.page_height=Inches(11.69)
    section.left_margin=Inches(.68);section.right_margin=Inches(.68)
    section.top_margin=Inches(.74);section.bottom_margin=Inches(.7)
    section.header_distance=Inches(.35);section.footer_distance=Inches(.35)
    normal=doc.styles["Normal"]
    set_font(normal,"Calibri",9,GREEN)
    normal.paragraph_format.space_after=Pt(4)
    for name,size in (("Title",26),("Heading 1",16),("Heading 2",11)):
        st=doc.styles[name]
        set_font(st,"Calibri",size,GREEN)
        st.paragraph_format.keep_with_next=True
        st.paragraph_format.space_before=Pt(9)
        st.paragraph_format.space_after=Pt(4)
    code=doc.styles.add_style("SourceCode",WD_STYLE_TYPE.PARAGRAPH)
    set_font(code,"Liberation Mono",7.25,RGBColor(38,47,42))
    code.paragraph_format.space_before=Pt(0)
    code.paragraph_format.space_after=Pt(0)
    code.paragraph_format.line_spacing=1.05
    code.paragraph_format.keep_together=False
    code.paragraph_format.widow_control=False
    head=section.header.paragraphs[0];head.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    rr=head.add_run("DABBERNA  /  COMPLETE SOURCE CODE");rr.font.size=Pt(8);rr.font.color.rgb=MUTED
    foot=section.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    rr=foot.add_run("Dabberna · Source reference · Page ");rr.font.size=Pt(8);rr.font.color.rgb=MUTED
    page_number(foot)

    doc.add_paragraph("DABBERNA  |  دبّرنا","Subtitle")
    doc.add_heading("Complete Source Code",0)
    p=doc.add_paragraph("Family spending, lifestyle and planning application")
    p.style="Subtitle"
    sha=os.environ.get("GITHUB_SHA") or subprocess.check_output(["git","rev-parse","HEAD"],text=True).strip()
    doc.add_paragraph("GitHub repository: https://github.com/fareshashemabdallah/wafferli-v13")
    doc.add_paragraph("Revision: "+sha)
    doc.add_paragraph("Generated: "+datetime.now().strftime("%d %B %Y"))
    doc.add_paragraph(str(len(files))+" tracked text files are reproduced in full. "+str(len(excluded))+" binary/unsupported files omitted.")
    doc.add_paragraph("This Word document is a searchable reference, not an executable application. Use the GitHub repository or ZIP for editing and building. Personal information saved in users' browsers is not included.")
    if excluded:doc.add_paragraph("Non-text files omitted: "+", ".join(excluded))
    doc.add_page_break()
    doc.add_heading("File index",1)
    current=None
    for name,_ in files:
        section_name=kind(name)
        if section_name!=current:
            doc.add_heading(section_name,2)
            current=section_name
        p=doc.add_paragraph(name)
        p.paragraph_format.space_after=Pt(1)
        for run in p.runs:run.font.size=Pt(7.8)
    doc.add_page_break()
    doc.add_heading("Source files",1)
    current=None
    for index,(name,content) in enumerate(files,1):
        section_name=kind(name)
        if section_name!=current:
            doc.add_heading(section_name,1)
            current=section_name
        doc.add_heading(f"{index:03d}   {name}",2)
        p=doc.add_paragraph(str(len(content.splitlines()))+" lines · UTF-8")
        p.paragraph_format.space_after=Pt(3)
        for run in p.runs:run.font.size=Pt(7);run.font.color.rgb=MUTED
        lines=content.replace("\r\n","\n").replace("\r","\n").split("\n")
        for offset in range(0,len(lines),22):
            para=doc.add_paragraph(style="SourceCode")
            para.add_run("\n".join(lines[offset:offset+22]))
    doc.core_properties.title="Dabberna – Complete Source Code"
    doc.core_properties.subject="All tracked text source files for the Dabberna family app"
    doc.save(OUT)
    verified=Document(OUT)
    headings=[p.text for p in verified.paragraphs if p.style.name=="Heading 2"]
    assert all(any(name in heading for heading in headings) for name,_ in files), "Incomplete file heading index"
    print(f"Created {OUT}, {OUT.stat().st_size} bytes, {len(files)} files")

if __name__=="__main__":
    main()
